"""
to_excel.py — 账本导出 Excel（分存储专用）
直接读取 ledger.json（整数分），输出元为单位的可读 Excel。
"""

import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

LEDGER = '/Users/labixiaoxin/.qclaw/workspace/data/ledger.json'
OUT    = '/Users/labixiaoxin/Desktop/4月账单_最新.xlsx'

def fen2yuan(fen):
    """分 → 元，浮点显示"""
    return fen / 100

def export():
    data = json.load(open(LEDGER))
    entries = sorted(data['entries'], key=lambda e: (e.get('date',''), e.get('created_at','')))

    wb = Workbook()
    ws = wb.active
    ws.title = "4月账单"

    # ── 样式 ──────────────────────────────────────────────────────
    hdr_font   = Font(bold=True, color="FFFFFF", size=11)
    hdr_fill   = PatternFill("solid", fgColor="4A4A4A")
    grn_fill   = PatternFill("solid", fgColor="E8F5E9")   # 预算内
    org_fill   = PatternFill("solid", fgColor="FFF3E0")   # 预算外
    blu_fill   = PatternFill("solid", fgColor="E3F2FD")   # 报销
    alt_fill   = PatternFill("solid", fgColor="F5F5F5")
    wht_fill   = PatternFill("solid", fgColor="FFFFFF")
    ylw_fill   = PatternFill("solid", fgColor="FFFDE7")   # 待确认
    thin = Side(style='thin', color="DDDDDD")
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    C = Alignment(horizontal="center", vertical="center")
    L = Alignment(horizontal="left",   vertical="center")
    R = Alignment(horizontal="right",  vertical="center")

    headers = ["日期","商户","类型","金额(元)","备注","预算","来源"]
    ws.append(headers)
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col)
        c.font = hdr_font; c.fill = hdr_fill
        c.alignment = C; c.border = bdr

    for idx, e in enumerate(entries):
        note = e.get('note','')
        ib   = e.get('in_budget', True)
        typ  = e.get('type','')
        src  = e.get('source','')
        amt_fen = e.get('amount', 0)
        pending  = e.get('pending', False)

        # 预算标签
        if '报销' in note and typ == '收入': bl = '报销收入'
        elif '报销' in note and typ == '支出': bl = '报销相抵'
        elif not ib: bl = '预算外'
        else:       bl = '✓ 预算内'

        row = idx + 2
        ws.cell(row=row,column=1,value=e.get('date','')).alignment = C
        ws.cell(row=row,column=2,value=e.get('merchant','')).alignment = L
        ws.cell(row=row,column=3,value=typ).alignment = C
        # 金额：分→元
        c_amt = ws.cell(row=row,column=4,value=fen2yuan(amt_fen))
        c_amt.number_format = '¥#,##0.00;(¥#,##0.00);"-"'
        c_amt.alignment = R
        ws.cell(row=row,column=5,value=note).alignment = L
        ws.cell(row=row,column=6,value=bl).alignment = C
        ws.cell(row=row,column=7,value=src).alignment = C   # 来源

        # 背景色
        if pending:        fill = ylw_fill  # 待确认
        elif '报销' in note: fill = blu_fill
        elif not ib and typ=='支出': fill = org_fill
        elif ib and typ=='支出':    fill = grn_fill
        else:                   fill = wht_fill if idx%2==0 else alt_fill

        for col in range(1,8):
            c = ws.cell(row=row,column=col)
            c.fill = fill; c.border = bdr

    ws.column_dimensions['A'].width = 13
    ws.column_dimensions['B'].width = 26
    ws.column_dimensions['C'].width = 8
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 12

    # ── 汇总 ──────────────────────────────────────────────────────
    in_items = [e for e in entries if e.get('type')=='支出' and e.get('in_budget',True) and e.get('note')!='已退款' and not e.get('pending')]
    out_items= [e for e in entries if e.get('type')=='支出' and not e.get('in_budget',True) and not e.get('pending')]
    reimb   = [e for e in entries if '报销' in e.get('note','') and e.get('type')=='收入']
    pending_items = [e for e in entries if e.get('pending')]

    in_sum  = sum(abs(e['amount']) for e in in_items)
    out_sum = sum(abs(e['amount']) for e in out_items)
    reimb_sum = sum(e['amount'] for e in reimb)
    remaining = 100000 - in_sum

    r = len(entries) + 3
    ws.cell(row=r,column=1,value="--- 4月汇总 ---").font = Font(bold=True,size=11)
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=3)
    for label, val, note, bold in [
        ("月度预算",       1000.00, "2026-04",   True),
        ("预算内支出",     fen2yuan(in_sum),     f"({len(in_items)}条)", False),
        ("预算外支出",     fen2yuan(out_sum),     f"({len(out_items)}条)", False),
        ("报销收入",       fen2yuan(reimb_sum),  f"({len(reimb)}条)",   False),
        ("💰 预算剩余",    fen2yuan(remaining),  "= 1000 - 预算内支出", True),
    ]:
        r+=1; ws.cell(row=r,column=1,value=label)
        c = ws.cell(row=r,column=4,value=val)
        c.number_format='¥#,##0.00;(¥#,##0.00);"-"'
        c.font=Font(bold=bold, color="0066CC" if "剩余" in label and remaining>=0 else "000000")
        ws.cell(row=r,column=5,value=note)

    if pending_items:
        r+=1; ws.cell(row=r,column=1,value=f"⚠️ 待确认条目: {len(pending_items)}条（黄底，需微信确认后入账）").font=Font(bold=True,color="E65100")

    # ── 图例 ──────────────────────────────────────────────────────
    r+=2
    for i,(lbl,clr) in enumerate([("✓ 预算内","E8F5E9"),("预算外","FFF3E0"),("报销","E3F2FD"),("⚠️ 待确认","FFFDE7")]):
        c = ws.cell(row=r,column=1+i*2,value=lbl)
        c.fill=PatternFill("solid",fgColor=clr); c.border=bdr; c.alignment=C

    wb.save(OUT)
    print(f"✅ 已导出: {OUT}")
    print(f"总条目: {len(entries)}条 | 预算内{in_sum/100:.2f} | 预算外{out_sum/100:.2f} | 剩余{remaining/100:.2f}")
    print(f"待确认: {len(pending_items)}条")

if __name__ == '__main__':
    export()
