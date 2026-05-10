#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
音频计费系统 v3 - 简洁清新版
"""

import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import subprocess
import os
from pathlib import Path
import re
from openpyxl import Workbook
from openpyxl.styles import Font
from datetime import datetime


class AudioBillingApp:
    def __init__(self, root):
        self.root = root
        self.root.title("音频计费系统")
        self.root.geometry("1000x750")
        self.root.configure(bg="#f5f7fa")
        
        # 数据
        self.scripts = []
        self.audios = []
        self.cv_map = {}  # CV映射表：原始名 -> 标准名
        self.default_price = 50.0
        self.special_prices = {}
        self.results = {}
        
        self.setup_ui()
    
    def setup_ui(self):
        # 标题
        header = tk.Frame(self.root, bg="#4A90D9", height=60)
        header.pack(fill=tk.X)
        header.pack_propagate(False)
        
        title = tk.Label(header, text="🎬 音频计费系统", 
                        font=("Arial", 20, "bold"), fg="white", bg="#4A90D9")
        title.pack(pady=15)
        
        # 主区域
        main = tk.Frame(self.root, bg="#f5f7fa")
        main.pack(fill=tk.BOTH, expand=True, padx=20, pady=15)
        
        # ===== 左侧 =====
        left = tk.Frame(main, bg="white", relief=tk.RAISED, bd=1)
        left.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        # 左侧标题
        tk.Label(left, text="📁 文件管理", font=("Arial", 12, "bold"), 
                bg="white", fg="#4A90D9").pack(anchor=tk.W, padx=15, pady=10)
        
        # 剧本区域
        script_frame = tk.Frame(left, bg="white")
        script_frame.pack(fill=tk.X, padx=15, pady=5)
        
        tk.Button(script_frame, text="📄 选择剧本文件", command=self.add_scripts,
                 bg="#4A90D9", fg="white", font=("Arial", 11, "bold"), width=18, height=2,
                 relief=tk.FLAT, cursor="hand2").pack(side=tk.LEFT, padx=2)
        
        self.script_count = tk.Label(script_frame, text="已选 0 个", 
                                   font=("Arial", 11, "bold"), bg="white", fg="#333")
        self.script_count.pack(side=tk.LEFT, padx=10)
        
        # 音频区域
        audio_frame = tk.Frame(left, bg="white")
        audio_frame.pack(fill=tk.BOTH, expand=True, padx=15, pady=10)
        
        tk.Button(audio_frame, text="🎵 选择音频文件", command=self.add_audios,
                 bg="#50C878", fg="white", font=("Arial", 11, "bold"), width=18, height=2,
                 relief=tk.FLAT, cursor="hand2").pack(pady=(0, 10))
        
        # 音频列表
        list_frame = tk.Frame(audio_frame, bg="#f5f7fa", relief=tk.SUNKEN, bd=1)
        list_frame.pack(fill=tk.BOTH, expand=True)
        
        scroll = tk.Scrollbar(list_frame)
        scroll.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.audio_list = tk.Listbox(list_frame, yscrollcommand=scroll.set,
                                     font=("Arial", 11), height=10,
                                     bg="#f5f7fa", relief=tk.FLAT, fg="#333")
        self.audio_list.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scroll.config(command=self.audio_list.yview)
        
        # 列表操作
        btn_row = tk.Frame(audio_frame, bg="white")
        btn_row.pack(fill=tk.X, pady=5)
        
        tk.Button(btn_row, text="删除选中", command=self.remove_audio,
                 bg="#FF6B6B", fg="white", font=("Arial", 10, "bold"), 
                 width=10, relief=tk.FLAT).pack(side=tk.LEFT, padx=2)
        tk.Button(btn_row, text="清空", command=self.clear_audios,
                 bg="#95A5A6", fg="white", font=("Arial", 10, "bold"),
                 width=8, relief=tk.FLAT).pack(side=tk.LEFT, padx=2)
        
        self.audio_count = tk.Label(btn_row, text="0 个音频",
                                   font=("Arial", 10, "bold"), bg="white", fg="#333")
        self.audio_count.pack(side=tk.RIGHT)
        
        # ===== 右侧 =====
        right = tk.Frame(main, bg="white", relief=tk.RAISED, bd=1, width=380)
        right.pack(side=tk.RIGHT, fill=tk.BOTH, padx=(15, 0))
        right.pack_propagate(False)
        
        # 右侧标题
        tk.Label(right, text="⚙️ 设置", font=("Arial", 12, "bold"), 
                bg="white", fg="#4A90D9").pack(anchor=tk.W, padx=15, pady=10)
        
        # CV映射表
        cv_label_frame = tk.LabelFrame(right, text="👤 CV映射表（防止一个CV多个名字）", 
                                       font=("Arial", 10, "bold"), bg="white", fg="#333")
        cv_label_frame.pack(fill=tk.X, padx=10, pady=5)
        
        tk.Label(cv_label_frame, text="格式: 原始名=标准名", 
                font=("Arial", 9), bg="white", fg="#666").pack(anchor=tk.W, padx=5)
        
        cv_frame = tk.Frame(cv_label_frame, bg="#fff", relief=tk.SUNKEN, bd=1)
        cv_frame.pack(fill=tk.X, padx=5, pady=5)
        
        self.cv_text = tk.Text(cv_frame, height=3, font=("Courier", 10), 
                              bg="#fff", fg="#000", relief=tk.FLAT)
        self.cv_text.pack(fill=tk.X, padx=2, pady=2)
        self.cv_text.insert(tk.END, "# 例: 哒小兜CV=哒小兜\n")
        
        # 间隔设置
        gap_frame = tk.Frame(right, bg="white")
        gap_frame.pack(fill=tk.X, padx=15, pady=8)
        
        tk.Label(gap_frame, text="句子间隔(秒):", font=("Arial", 11, "bold"), 
                bg="white", fg="#333").pack(side=tk.LEFT)
        self.gap_var = tk.StringVar(value="0.5")
        tk.Entry(gap_frame, textvariable=self.gap_var, width=8, 
                font=("Arial", 12), fg="#000", bg="#fff").pack(side=tk.LEFT, padx=5)
        
        # 默认单价
        price_frame = tk.Frame(right, bg="white")
        price_frame.pack(fill=tk.X, padx=15, pady=8)
        
        tk.Label(price_frame, text="默认单价(元/分钟):", font=("Arial", 11, "bold"), 
                bg="white", fg="#333").pack(side=tk.LEFT)
        self.price_var = tk.StringVar(value="50")
        tk.Entry(price_frame, textvariable=self.price_var, width=8,
                font=("Arial", 12), fg="#000", bg="#fff").pack(side=tk.LEFT, padx=5)
        
        # 角色单价列表
        tk.Label(right, text="特殊角色单价 (格式: 角色:价格)", 
                font=("Arial", 10, "bold"), bg="white", fg="#333").pack(anchor=tk.W, padx=15, pady=(10, 5))
        
        price_list_frame = tk.Frame(right, bg="#f5f7fa", relief=tk.SUNKEN, bd=1)
        price_list_frame.pack(fill=tk.BOTH, expand=True, padx=15, pady=5)
        
        scroll2 = tk.Scrollbar(price_list_frame)
        scroll2.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.price_text = tk.Text(price_list_frame, yscrollcommand=scroll2.set,
                                  font=("Courier", 11), height=6, bg="#fff", fg="#000", relief=tk.FLAT)
        self.price_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scroll2.config(command=self.price_text.yview)
        
        self.price_text.insert(tk.END, "# 大部分角色用默认单价\n# 只有特殊的才写:\n# 朵拉:60\n# 莉莉娜:55\n")
        
        # ===== 底部 =====
        bottom = tk.Frame(self.root, bg="#fff", height=80)
        bottom.pack(fill=tk.X, side=tk.BOTTOM)
        
        # 开始按钮
        tk.Button(bottom, text="▶️ 开始计费", command=self.run_billing,
                 bg="#4A90D9", fg="white", font=("Arial", 14, "bold"),
                 width=15, height=2, relief=tk.FLAT, cursor="hand2").pack(side=tk.LEFT, padx=20, pady=15)
        
        # 导出按钮
        tk.Button(bottom, text="📊 导出Excel", command=self.export_excel,
                 bg="#50C878", fg="white", font=("Arial", 12, "bold"),
                 width=12, height=2, relief=tk.FLAT, cursor="hand2").pack(side=tk.LEFT, padx=10, pady=15)
        
        # 状态
        self.status = tk.Label(bottom, text="就绪", font=("Arial", 11, "bold"), bg="#fff", fg="#333")
        self.status.pack(side=tk.RIGHT, padx=20)
    
    # ========== 功能 ==========
    def add_scripts(self):
        files = filedialog.askopenfilenames(
            title="选择剧本文件",
            filetypes=[("Word文档", "*.docx"), ("所有文件", "*")]
        )
        self.scripts.extend(files)
        self.script_count.config(text=f"已选 {len(self.scripts)} 个")
    
    def add_audios(self):
        files = filedialog.askopenfilenames(
            title="选择音频文件",
            filetypes=[("音频文件", "*.mp3 *.wav *.m4a"), ("所有文件", "*")]
        )
        for f in files:
            self.audios.append(f)
            name = Path(f).name
            # 提取演员名
            actor = self.extract_actor(name)
            self.audio_list.insert(tk.END, f"🎵 {name} → {actor}")
        self.audio_count.config(text=f"{len(self.audios)} 个音频")
    
    def extract_actor(self, filename):
        """从文件名提取演员名，并应用CV映射表"""
        # 去掉扩展名
        name = Path(filename).stem
        
        # 按常见分隔符分割
        parts = re.split(r'[_\-\s]+', name)
        
        # 尝试找出演员名（通常在最后或倒数第二个）
        for part in reversed(parts):
            if len(part) >= 2 and not re.match(r'^\d+$', part):  # 不是纯数字
                # 检查CV映射表
                if part in self.cv_map:
                    return self.cv_map[part]
                return part
        
        # 没找到，返回整个文件名
        return name
    
    def parse_cv_map(self):
        """解析CV映射表"""
        self.cv_map = {}
        for line in self.cv_text.get("1.0", tk.END).strip().split('\n'):
            if '=' in line and not line.startswith('#'):
                try:
                    k, v = line.split('=', 1)
                    self.cv_map[k.strip()] = v.strip()
                except:
                    pass
    
    def remove_audio(self):
        sel = self.audio_list.curselection()
        if sel:
            self.audio_list.delete(sel)
            self.audios.pop(sel[0])
            self.audio_count.config(text=f"{len(self.audios)} 个音频")
    
    def clear_audios(self):
        self.audios = []
        self.audio_list.delete(0, tk.END)
        self.audio_count.config(text="0 个音频")
    
    def parse_prices(self):
        self.default_price = float(self.price_var.get())
        self.special_prices = {}
        for line in self.price_text.get("1.0", tk.END).strip().split('\n'):
            if ':' in line and not line.startswith('#'):
                try:
                    role, price = line.split(':')
                    self.special_prices[role.strip()] = float(price.strip())
                except:
                    pass
    
    def run_billing(self):
        if not self.scripts:
            messagebox.showerror("错误", "请先选择剧本文件")
            return
        if not self.audios:
            messagebox.showerror("错误", "请先添加音频文件")
            return
        
        self.parse_cv_map()  # 解析CV映射表
        self.parse_prices()
        
        self.status.config(text="处理中...")
        self.root.update()
        
        # 进度条窗口
        pw = tk.Toplevel(self.root)
        pw.title("处理进度")
        pw.geometry("400x120")
        pw.resizable(False, False)
        
        tk.Label(pw, text="正在处理...", font=("Arial", 11, "bold"), fg="#333").pack(pady=8)
        
        pvar = tk.DoubleVar()
        ttk.Progressbar(pw, variable=pvar, maximum=100, length=350, mode="determinate").pack(pady=5)
        
        status = tk.Label(pw, text="", font=("Arial", 10), fg="#333")
        status.pack()
        
        self.results = {}
        total = len(self.audios)
        
        for i, audio in enumerate(self.audios, 1):
            pvar.set((i / total) * 100)
            status.config(text=f"{i}/{total}")
            pw.update()
            
            # 提取演员名
            name = Path(audio).name
            actor = self.extract_actor(name)
            
            try:
                # 用第一个剧本处理
                cmd = [
                    "python3",
                    "/Users/labixiaoxin/.qclaw/workspace/audio_billing/audio_billing.py",
                    self.scripts[0], audio,
                    "--gap", self.gap_var.get(),
                    "--model", "base",
                    "--default-price", self.price_var.get(),
                ]
                for r, p in self.special_prices.items():
                    cmd.extend(["--price", f"{r}:{p}"])
                
                result = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
                
                if result.returncode == 0:
                    parsed = self.parse_result(result.stdout)
                    if parsed:
                        self.results[actor] = parsed
            except:
                pass
        
        pvar.set(100)
        pw.update()
        pw.after(500, pw.destroy)  # 0.5秒后自动关闭
        
        total_cost = sum(sum(d.get('cost', 0) for d in r.values()) for r in self.results.values())
        self.status.config(text=f"完成! 共 {len(self.results)} 人, 总计 {total_cost:.2f} 元")
        
        messagebox.showinfo("完成", f"处理完成!\n共 {len(self.results)} 个演员\n总计: {total_cost:.2f} 元\n\n点击「导出Excel」保存报告")
    
    def parse_result(self, output):
        result = {}
        pattern = r'([^\n:]+?):\s+说话时长:.*?费用:\s*[\d.]+\s*×\s*([\d.]+)元/分钟\s*=\s*([\d.]+)元'
        for m in re.findall(pattern, output, re.DOTALL):
            if m[0] and m[0] != "未匹配":
                result[m[0]] = {'price': float(m[1]), 'cost': float(m[2])}
        return result
    
    def export_excel(self):
        if not self.results:
            messagebox.showinfo("提示", "暂无数据，请先运行计费")
            return
        
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "计费报告"
            
            ws['A1'] = "音频计费报告"
            ws['A1'].font = Font(size=14, bold=True)
            ws['A2'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            
            headers = ["演员", "角色", "费用(元)"]
            for col, h in enumerate(headers, 1):
                ws.cell(row=4, column=col, value=h).font = Font(bold=True)
            
            row = 5
            total = 0
            for actor, roles in self.results.items():
                for role, data in roles.items():
                    ws.cell(row=row, column=1, value=actor)
                    ws.cell(row=row, column=2, value=role)
                    ws.cell(row=row, column=3, value=round(data['cost'], 2))
                    row += 1
            
            ws.cell(row=row, column=1, value="总计").font = Font(bold=True)
            ws.cell(row=row, column=3, value=round(sum(sum(d['cost'] for d in r.values()) for r in self.results.values()), 2))
            ws.cell(row=row, column=3).font = Font(bold=True)
            
            for col in 'ABC':
                ws.column_dimensions[col].width = 15
            
            path = os.path.expanduser("~/Desktop/计费报告.xlsx")
            wb.save(path)
            os.system(f"open {path}")
        except Exception as e:
            messagebox.showerror("错误", f"导出失败: {str(e)}")


if __name__ == "__main__":
    root = tk.Tk()
    app = AudioBillingApp(root)
    root.mainloop()